# -*- coding: utf-8 -*-
"""
电费结算单提取与比对通用工具
================================

封装两个核心方法，并以通用工具类 :class:`BillSettlementTool` 对外提供：

方法 1  ``filter_bill_files``
    过滤出所有「电费结算单」文件，并去掉内容完全相同（字节级）的重复文件。
    注意：结算单必须是「电费结算单」，不能是「电量结算单」（以及「核算单 /
    补贴核算单」等其它单据一律排除）。

方法 2  ``extract_records``
    从方法 1 得到的文件中，提取三个数据项：

        - 编号      —— 单据表头「电厂（交易对象）编号」
        - 购电月份  —— 单据「购电月份」字段（如 2026 年 01 月）
        - 电价      —— 单据「电价」列。该列通常有多行（市场化电费、机制电价差、
                       结算小计…），按要求 **取最后一行（结算小计）的原始值**；
                       一行都没有时用 ``0`` 代替。

    一个文件可能包含多条记录（例如华尔特的 PDF 一份 13 页 = 13 个编号），
    每页 / 每张图片各产出一条记录。

比对
    ``match_with_excel`` 把提取结果与 Excel 做比对：

        - Excel 中「备注」为「新丰」的行不参与比对（去掉新丰，剩余共 76 条）；
        - 编号、购电月份、电价 **三项完全相等** 才算匹配成功（是相等不是约等于）；
        - 提取的数据可以多，但不能少：Excel 的 76 条必须全部被命中才算成功。

OCR 引擎
    使用 RapidOCR（基于 PP-OCR，模型随 pip 包一起分发，无需联网下载），
    对干净的电子单据识别准确率很高。若环境允许联网下载模型，也可改用 PaddleOCR。

依赖::

    pip install rapidocr-onnxruntime onnxruntime pymupdf pillow numpy openpyxl

命令行用法::

    python bill_settlement_tool.py <文件根目录> <对比用Excel.xlsx>
"""

from __future__ import annotations

import os
import re
import hashlib
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple

import numpy as np

# ----------------------------------------------------------------------------- #
# 数据结构
# ----------------------------------------------------------------------------- #

YearMonth = Tuple[int, int]  # (年, 月)


@dataclass
class BillRecord:
    """从单据某一页提取出来的一条记录。"""

    bill_id: Optional[str]          # 编号（电厂（交易对象）编号）
    month: Optional[YearMonth]      # 购电月份 -> (年, 月)
    price: float                    # 电价（电价列最后一行的值，没有则 0）
    source_file: str = ""           # 来源文件
    page: int = 1                   # 来源页码（PDF 多页时有用）
    title_ok: bool = True           # 该页标题是否确为「电费结算单」

    def compare_key(self) -> Tuple[str, YearMonth, float]:
        """用于精确比对的归一化键：(编号, (年,月), 电价保留 6 位)。"""
        return (
            (self.bill_id or "").strip(),
            self.month,
            round(float(self.price), 6),
        )


@dataclass
class ExcelRow:
    """Excel 中一行（已折算好的对比基准）。"""

    bill_id: str
    month: YearMonth
    price: float
    note: str = ""

    def compare_key(self) -> Tuple[str, YearMonth, float]:
        return (self.bill_id.strip(), self.month, round(float(self.price), 6))


@dataclass
class RowDetail:
    """单条 Excel 行与提取数据的逐项对比明细。"""

    excel: ExcelRow                       # Excel 基准行
    got_id: Optional[str]                 # 提取到的编号（用于对比的那条记录）
    got_month: Optional[YearMonth]        # 提取到的购电月份
    got_price: Optional[float]            # 提取到的电价
    id_ok: bool                           # 编号是否相等
    month_ok: bool                        # 购电月份是否相等
    price_ok: bool                        # 电价是否相等
    source: str = ""                      # 提取来源（文件:页）

    @property
    def matched(self) -> bool:
        return self.id_ok and self.month_ok and self.price_ok

    @staticmethod
    def _ym(ym: Optional[YearMonth]) -> str:
        return f"{ym[0]}-{ym[1]:02d}" if ym else "—"

    def line(self) -> str:
        mark = lambda ok: "✅" if ok else "❌"
        gid = self.got_id or "—"
        gprice = "—" if self.got_price is None else f"{self.got_price:.6f}".rstrip("0").rstrip(".")
        eprice = f"{self.excel.price:.6f}".rstrip("0").rstrip(".")
        return (
            f"[{'✅' if self.matched else '❌'}] "
            f"编号 {self.excel.bill_id}{mark(self.id_ok)}={gid}  "
            f"购电月份 {self._ym(self.excel.month)}{mark(self.month_ok)}={self._ym(self.got_month)}  "
            f"电价 {eprice}{mark(self.price_ok)}={gprice}  "
            f"[{self.excel.note}] {self.source}"
        )


@dataclass
class MatchResult:
    """比对结果汇总。"""

    total_excel: int
    matched: int
    unmatched_rows: List[ExcelRow] = field(default_factory=list)
    extracted_records: List[BillRecord] = field(default_factory=list)
    details: List[RowDetail] = field(default_factory=list)

    @property
    def success(self) -> bool:
        """Excel 全部 76 条都命中才算成功。"""
        return self.total_excel > 0 and self.matched == self.total_excel

    def summary(self) -> str:
        head = (
            f"匹配结果：{self.matched}/{self.total_excel} "
            f"-> {'成功 ✅' if self.success else '失败 ❌'}"
        )
        if self.unmatched_rows:
            lines = ["未匹配的 Excel 行："]
            for r in self.unmatched_rows:
                lines.append(
                    f"  编号={r.bill_id} 购电月份={r.month[0]}-{r.month[1]:02d} "
                    f"电价={r.price} 备注={r.note}"
                )
            head += "\n" + "\n".join(lines)
        return head

    def detail_report(self, only_fail: bool = False) -> str:
        """逐项对比明细：把 Excel 每一条与提取值并排打印，三项各标 ✅/❌。

        :param only_fail: 仅打印未匹配的行（默认 False，全部打印）。
        """
        rows = [d for d in self.details if (not only_fail or not d.matched)]
        lines = [f"逐项对比明细（共 {len(self.details)} 条，显示 {len(rows)} 条）："]
        for i, d in enumerate(rows, 1):
            lines.append(f"{i:>3}. {d.line()}")
        lines.append("")
        lines.append(self.summary())
        return "\n".join(lines)


# ----------------------------------------------------------------------------- #
# 工具类
# ----------------------------------------------------------------------------- #

class BillSettlementTool:
    """电费结算单 过滤 / 提取 / 比对 通用工具类。"""

    # —— 文件名判定 —— #
    # 是电费结算单：文件名含「结算单」且不含「电量」（从而排除「电量结算单」），
    # 其它如「核算单 / 分布式光伏发电补贴核算单」本身不含「结算单」，自然被排除。
    _KEEP_KEYWORD = "结算单"
    _EXCLUDE_KEYWORD = "电量"

    # —— 字段识别用正则 —— #
    _MONTH_RE = re.compile(r"(20\d{2})\D+?(\d{1,2})\D*月")   # 购电月份 2026年01月
    _ID_RE = re.compile(r"^\d{14,18}$")                       # 编号：14~18 位纯数字
    _PRICE_RE = re.compile(r"^\d+\.\d{4,}$")                  # 电价：小数点后≥4 位

    def __init__(self, dpi: int = 300, ocr_engine=None):
        """
        :param dpi:        PDF 渲染分辨率，越高识别越准但越慢，300 足够。
        :param ocr_engine: 可注入自定义 OCR；默认惰性创建 RapidOCR。
        """
        self.dpi = dpi
        self._ocr = ocr_engine

    # ------------------------------------------------------------------ #
    # OCR 引擎（惰性加载）
    # ------------------------------------------------------------------ #
    @property
    def ocr(self):
        if self._ocr is None:
            from rapidocr_onnxruntime import RapidOCR
            self._ocr = RapidOCR()
        return self._ocr

    # ================================================================== #
    # 方法 1：过滤电费结算单文件 + 去重
    # ================================================================== #
    @classmethod
    def is_bill_file(cls, filename: str) -> bool:
        """文件名是否属于「电费结算单」（排除电量结算单 / 核算单等）。"""
        name = os.path.basename(filename)
        return (cls._KEEP_KEYWORD in name) and (cls._EXCLUDE_KEYWORD not in name)

    @classmethod
    def filter_bill_files(cls, root_dir: str, verbose: bool = False) -> List[str]:
        """方法 1：递归过滤所有电费结算单文件，并去掉内容完全相同的重复文件。

        :param root_dir: 待扫描的根目录。
        :param verbose:  为 True 时打印命中 / 去重过程。
        :return: 去重后的电费结算单文件路径列表（按路径排序、内容唯一）。
        """
        seen_hashes: Dict[str, str] = {}
        result: List[str] = []
        dup = 0
        for dir_path, _dirs, files in os.walk(root_dir):
            for fname in sorted(files):
                if not cls.is_bill_file(fname):
                    continue
                full = os.path.join(dir_path, fname)
                digest = cls._file_sha256(full)
                if digest in seen_hashes:
                    dup += 1
                    if verbose:
                        print(f"  跳过(内容重复) {os.path.relpath(full, root_dir)}")
                    continue          # 内容相同的文件，跳过
                seen_hashes[digest] = full
                result.append(full)
                if verbose:
                    print(f"  命中 {os.path.relpath(full, root_dir)}")
        result.sort()
        if verbose:
            print(f"-- 命中 {len(result)} 个，去重跳过 {dup} 个 --\n")
        return result

    @staticmethod
    def _file_sha256(path: str, chunk: int = 1 << 20) -> str:
        h = hashlib.sha256()
        with open(path, "rb") as fp:
            while True:
                block = fp.read(chunk)
                if not block:
                    break
                h.update(block)
        return h.hexdigest()

    # ================================================================== #
    # 方法 2：从文件中提取 编号 / 购电月份 / 电价
    # ================================================================== #
    def extract_file(self, path: str, verbose: bool = False) -> List[BillRecord]:
        """从单个文件提取数据。

        一个文件可能有多页（如华尔特 PDF 一份 13 页 = 13 条记录），
        因此返回的是该文件的记录列表；单页文件则只有一条。

        :param path:    单个电费结算单文件路径（PDF / PNG / JPG）。
        :param verbose: 为 True 时打印该文件每一页的提取结果。
        :return: 该文件的 BillRecord 列表。
        """
        records: List[BillRecord] = []
        images = self._render_pages(path)
        if verbose:
            print(f"解析 {os.path.basename(path)}  （{len(images)} 页）")
        for page_idx, image in enumerate(images):
            tokens = self._ocr_tokens(image)
            rec = self._parse_page(tokens)
            rec.source_file = path
            rec.page = page_idx + 1
            records.append(rec)
            if verbose:
                ym = f"{rec.month[0]}-{rec.month[1]:02d}" if rec.month else "—"
                print(
                    f"      p{rec.page}: 编号={rec.bill_id or '—'}  "
                    f"购电月份={ym}  电价={rec.price}"
                )
        return records

    def extract_records(self, files: List[str], verbose: bool = False) -> List[BillRecord]:
        """方法 2：从方法 1 得到的（多个）文件中提取数据。

        内部对每个文件调用 :meth:`extract_file`。

        :param files:   方法 1 返回的文件列表。
        :param verbose: 为 True 时打印每个文件 / 每一页的提取过程。
        :return: BillRecord 列表。
        """
        records: List[BillRecord] = []
        total = len(files)
        for fi, path in enumerate(files, 1):
            if verbose:
                print(f"[{fi}/{total}] ", end="")
            records.extend(self.extract_file(path, verbose=verbose))
        if verbose:
            print(f"-- 提取完成，共 {len(records)} 条记录 --\n")
        return records

    def extract_from_dir(self, root_dir: str) -> List[BillRecord]:
        """便捷方法：方法 1 + 方法 2 一步到位。"""
        return self.extract_records(self.filter_bill_files(root_dir))

    # ---- 渲染：PDF 多页 / 图片单页 -> RGB numpy 数组 ---- #
    def _render_pages(self, path: str) -> List[np.ndarray]:
        ext = os.path.splitext(path)[1].lower()
        images: List[np.ndarray] = []
        if ext == ".pdf":
            import fitz  # PyMuPDF
            doc = fitz.open(path)
            for page in doc:
                pix = page.get_pixmap(dpi=self.dpi)
                arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                    pix.height, pix.width, pix.n
                )
                if pix.n >= 3:
                    arr = arr[:, :, :3].copy()
                else:  # 灰度 -> RGB
                    arr = np.repeat(arr, 3, axis=2)
                images.append(arr)
            doc.close()
        else:  # png / jpg / jpeg ...
            from PIL import Image
            images.append(np.array(Image.open(path).convert("RGB")))
        return images

    # ---- OCR：图片 -> 带坐标的 token 列表 ---- #
    def _ocr_tokens(self, image: np.ndarray) -> List[dict]:
        result, _ = self.ocr(image)
        tokens: List[dict] = []
        if not result:
            return tokens
        for box, text, conf in result:
            xs = [p[0] for p in box]
            ys = [p[1] for p in box]
            tokens.append(
                dict(
                    text=text.strip(),
                    xc=sum(xs) / 4.0,
                    yc=sum(ys) / 4.0,
                    conf=conf,
                )
            )
        return tokens

    # ---- 解析单页 token -> BillRecord ---- #
    def _parse_page(self, tokens: List[dict]) -> BillRecord:
        title_ok = any(t["text"] == "电费结算单" for t in tokens)
        return BillRecord(
            bill_id=self._find_id(tokens),
            month=self._find_month(tokens),
            price=self._find_price(tokens),
            title_ok=title_ok,
        )

    def _find_month(self, tokens: List[dict]) -> Optional[YearMonth]:
        """购电月份：匹配 20xx年x月（排除含「日」的截表时间）。"""
        for t in tokens:
            txt = t["text"].replace(" ", "")
            if "日" in txt:
                continue
            m = self._MONTH_RE.search(txt)
            if m and txt.endswith("月"):
                return (int(m.group(1)), int(m.group(2)))
        # 兜底：放宽到任何含「年..月」的 token
        for t in tokens:
            m = self._MONTH_RE.search(t["text"].replace(" ", ""))
            if m and "日" not in t["text"]:
                return (int(m.group(1)), int(m.group(2)))
        return None

    def _find_id(self, tokens: List[dict]) -> Optional[str]:
        """编号：定位「电厂（交易对象）编号」表头，取其正下方最近的纯数字串。"""
        header = None
        for t in tokens:
            if "编号" in t["text"] and ("电厂" in t["text"] or "交易对象" in t["text"]):
                header = t
                break
        ids = [t for t in tokens if self._ID_RE.match(t["text"])]
        if not ids:
            return None
        if header is not None:
            below = [t for t in ids if t["yc"] > header["yc"]] or ids
            # 与表头水平位置最接近、且尽量靠上的那个
            best = min(
                below,
                key=lambda t: abs(t["xc"] - header["xc"]) + 0.001 * abs(t["yc"] - header["yc"]),
            )
            return best["text"]
        # 没识别到表头时，取最长的数字串
        return max(ids, key=lambda t: len(t["text"]))["text"]

    def _find_price(self, tokens: List[dict]) -> float:
        """电价：电价列里小数点后≥4 位的数值，取最后一行（纵坐标最大）的值。"""
        prices = [t for t in tokens if self._PRICE_RE.match(t["text"])]
        if not prices:
            return 0.0  # 一个都没有 -> 用 0 代替
        last = max(prices, key=lambda t: t["yc"])  # 取最后一个（结算小计行）
        return float(last["text"])

    # ================================================================== #
    # 读取 Excel + 比对
    # ================================================================== #
    @staticmethod
    def load_excel(
        xlsx_path: str,
        sheet: Optional[str] = None,
        drop_notes: Tuple[str, ...] = ("新丰",),
    ) -> List[ExcelRow]:
        """读取对比基准 Excel。

        约定列顺序：序号 | 电厂（交易对象）编号 | 购电月份 | 电价 | 备注。
        「购电月份」可能是 Excel 日期序列号（如 46023）或日期，统一折算成 (年, 月)。
        备注命中 ``drop_notes``（默认「新丰」）的行被剔除。
        """
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows: List[ExcelRow] = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if raw is None or len(raw) < 4:
                continue
            _seq, bill_id, month, price = raw[0], raw[1], raw[2], raw[3]
            note = str(raw[4]).strip() if len(raw) > 4 and raw[4] is not None else ""
            if bill_id is None or price is None or month is None:
                continue
            if any(d in note for d in drop_notes):
                continue  # 去掉新丰
            rows.append(
                ExcelRow(
                    bill_id=str(bill_id).strip(),
                    month=BillSettlementTool._to_year_month(month),
                    price=round(float(price), 6),
                    note=note,
                )
            )
        return rows

    @staticmethod
    def _to_year_month(value) -> YearMonth:
        """把 Excel 单元格的月份值折算成 (年, 月)。"""
        if isinstance(value, datetime):
            return (value.year, value.month)
        if isinstance(value, (int, float)):
            d = datetime(1899, 12, 30) + timedelta(days=int(value))
            return (d.year, d.month)
        # 文本形式，如 "2026年01月" / "2026-01"
        m = re.search(r"(20\d{2})\D+?(\d{1,2})", str(value))
        if m:
            return (int(m.group(1)), int(m.group(2)))
        raise ValueError(f"无法解析购电月份: {value!r}")

    def match_with_excel(
        self,
        records: List[BillRecord],
        xlsx_path: str,
        sheet: Optional[str] = None,
    ) -> MatchResult:
        """把提取记录与 Excel 比对（三项完全相等才算命中）。

        Excel 的每一条都必须被提取数据命中；提取数据允许有多余。
        """
        excel_rows = self.load_excel(xlsx_path, sheet=sheet)

        valid = [r for r in records if r.bill_id and r.month is not None]
        # 三项全等的精确索引
        exact_keys = {r.compare_key() for r in valid}
        # 退化索引：用于在未命中时挑一条最接近的记录展示逐项差异
        by_id_month: Dict[Tuple[str, YearMonth], BillRecord] = {}
        by_id: Dict[str, BillRecord] = {}
        for r in valid:
            by_id_month.setdefault((r.bill_id.strip(), r.month), r)
            by_id.setdefault(r.bill_id.strip(), r)

        matched = 0
        unmatched: List[ExcelRow] = []
        details: List[RowDetail] = []
        for row in excel_rows:
            if row.compare_key() in exact_keys:
                matched += 1
                # 命中：取同 (编号,月份) 的记录展示其值
                rec = by_id_month.get((row.bill_id, row.month))
                details.append(self._make_detail(row, rec, full_match=True))
            else:
                unmatched.append(row)
                # 未命中：优先按 (编号,月份) 找，再按 编号 找，作为对照展示
                rec = by_id_month.get((row.bill_id, row.month)) or by_id.get(row.bill_id)
                details.append(self._make_detail(row, rec, full_match=False))

        return MatchResult(
            total_excel=len(excel_rows),
            matched=matched,
            unmatched_rows=unmatched,
            extracted_records=records,
            details=details,
        )

    @staticmethod
    def _make_detail(row: ExcelRow, rec: Optional[BillRecord], full_match: bool) -> RowDetail:
        if rec is None:
            return RowDetail(row, None, None, None, False, False, False, "未提取到对应编号")
        id_ok = (rec.bill_id or "").strip() == row.bill_id.strip()
        month_ok = rec.month == row.month
        price_ok = round(float(rec.price), 6) == row.price
        src = f"{os.path.basename(rec.source_file)}:p{rec.page}"
        return RowDetail(
            row, rec.bill_id, rec.month, round(float(rec.price), 6),
            id_ok, month_ok, price_ok, src,
        )

    # ================================================================== #
    # 一站式流程
    # ================================================================== #
    def run(self, root_dir: str, xlsx_path: str, sheet: Optional[str] = None) -> MatchResult:
        """方法1 -> 方法2 -> 比对，一条龙。"""
        files = self.filter_bill_files(root_dir)
        records = self.extract_records(files)
        return self.match_with_excel(records, xlsx_path, sheet=sheet)


# ----------------------------------------------------------------------------- #
# 命令行入口
# ----------------------------------------------------------------------------- #
def _main(argv: List[str]) -> int:
    args = [a for a in argv[1:] if not a.startswith("-")]
    flags = {a for a in argv[1:] if a.startswith("-")}
    if len(args) < 2:
        print(__doc__)
        print(
            "用法: python bill_settlement_tool.py <文件根目录> <对比用Excel.xlsx> "
            "[sheet名] [--detail] [--fail-only]"
        )
        return 1
    root_dir, xlsx_path = args[0], args[1]
    sheet = args[2] if len(args) > 2 else None

    # 默认打印过程；--quiet 关闭
    verbose = "--quiet" not in flags

    tool = BillSettlementTool(dpi=300)

    print("[方法1] 过滤电费结算单文件（排除电量结算单）并按内容去重：")
    files = tool.filter_bill_files(root_dir, verbose=verbose)
    print(f"[方法1] 过滤+去重后电费结算单文件: {len(files)} 个\n")

    print("[方法2] OCR 提取 编号/购电月份/电价（PDF 多页较慢，请稍候）：")
    records = tool.extract_records(files, verbose=verbose)
    print(f"[方法2] 共提取记录: {len(records)} 条\n")

    result = tool.match_with_excel(records, xlsx_path, sheet=sheet)
    # 默认打印逐项对比明细；--fail-only 只看未匹配的
    if "--fail-only" in flags:
        print(result.detail_report(only_fail=True))
    else:
        print(result.detail_report(only_fail=False))
    return 0 if result.success else 2


if __name__ == "__main__":
    import sys

    raise SystemExit(_main(sys.argv))
